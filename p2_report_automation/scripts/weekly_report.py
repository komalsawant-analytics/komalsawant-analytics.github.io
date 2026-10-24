"""
Weekly Store Sales & Stock Report — on autopilot
=================================================
Sample project · simulated data for a fictional 4-store home & kitchen retailer ("Kora Home").

Every Monday each store sends its POS export in its OWN format (xlsx with a banner,
csv, semicolon csv, lowercase codes, text amounts, three date styles). This script:
  1. reads every file in the week's folder using a per-store mapping (no manual copy-paste),
  2. cleans and unifies them, and compares with the previous week,
  3. writes a formatted Excel report with live formulas, a chart and a reorder-alert sheet,
  4. writes a short e-mail summary (HTML) ready to send to the owner.

Schedule it with Windows Task Scheduler / cron, or port the same logic to Google Apps Script.

Run:  python weekly_report.py --week 2026-10-12
"""
import argparse
from pathlib import Path
from datetime import datetime, timedelta
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

BASE = Path(__file__).resolve().parent.parent
IN, OUT = BASE / "input", BASE / "output"
OUT.mkdir(exist_ok=True)

# ---- One entry per store: how to read its export. Adding a store = adding a block here.
STORES = {
    "Andheri": dict(file="Andheri_POS.xlsx", read=dict(), date="Date", date_fmt=None,
                    code="Item Code", name=None, qty="Qty", amount="Amount"),
    "Bandra":  dict(file="bandra_sales.csv", read=dict(), date="bill_date", date_fmt="%Y-%m-%d",
                    code="sku", name=None, qty="qty_sold", amount="net_amt"),
    "Thane":   dict(file="THANE store report.xlsx", read=dict(skiprows=3), date="Sale Dt", date_fmt="%d/%m/%Y",
                    code=None, name="Product", qty="Pcs", amount="Value (Rs.)"),
    "Pune":    dict(file="pune-export.csv", read=dict(sep=";"), date="Datum", date_fmt="%d.%m.%Y",
                    code="Code", name=None, qty="Quantity", amount="Total"),
}
REORDER_DAYS = 7          # flag items with less than this many days of stock left
LOG = []

def read_store(folder, store, cfg, items):
    f = folder / cfg["file"]
    df = pd.read_excel(f, **cfg["read"]) if f.suffix == ".xlsx" else pd.read_csv(f, **cfg["read"])
    out = pd.DataFrame()
    out["date"] = pd.to_datetime(df[cfg["date"]], format=cfg["date_fmt"]).dt.date
    if cfg["code"]:
        out["item_code"] = df[cfg["code"]].astype(str).str.strip().str.upper()
    else:  # store only exports item names -> look up the code from the item master
        lookup = dict(zip(items["item_name"].str.upper().str.strip(), items["item_code"]))
        out["item_code"] = df[cfg["name"]].astype(str).str.upper().str.strip().map(lookup)
        LOG.append((store, "Export has item names only (upper-case), no codes", len(df), "Mapped names to codes via item master"))
    out["qty"] = pd.to_numeric(df[cfg["qty"]], errors="coerce")
    out["amount"] = pd.to_numeric(df[cfg["amount"]].astype(str).str.replace(",", ""), errors="coerce")
    out["store"] = store
    bad = out[["item_code", "qty", "amount"]].isna().any(axis=1)
    if bad.any():
        LOG.append((store, "Unreadable rows", bad.sum(), "Excluded - listed for store manager"))
    return out[~bad]

def load_week(week):
    folder = IN / f"week_{week}"
    items = pd.read_csv(IN / "item_master.csv")
    df = pd.concat([read_store(folder, s, c, items) for s, c in STORES.items()], ignore_index=True)
    df = df.merge(items, on="item_code", how="left")
    df["week"] = week
    return df

def main(week):
    prev_week = (datetime.strptime(week, "%Y-%m-%d") - timedelta(days=7)).strftime("%Y-%m-%d")
    LOG.extend([
        ("Bandra", "Item codes exported in lower case", 0, "Upper-cased"),
        ("Thane", "3 banner rows above the header; amounts as '1,299' text", 0, "Skipped banner; converted to numbers"),
        ("Pune", "Semicolon-separated CSV with DD.MM.YYYY dates", 0, "Read with ';' and explicit date format"),
    ])
    cur, prev = load_week(week), load_week(prev_week)
    data = pd.concat([prev, cur], ignore_index=True).sort_values(["week", "store", "date"])
    stock = pd.read_excel(IN / f"week_{week}" / "stock_on_hand.xlsx")
    stock.columns = ["store", "item_code", "closing_stock"]

    wb = Workbook()
    F = "Arial"
    hdr_fill, hdr_font = PatternFill("solid", fgColor="1F3A5F"), Font(name=F, bold=True, color="FFFFFF")
    thin = Border(bottom=Side(style="thin", color="D0D5DD"))
    def style_header(ws, row, ncol):
        for c in range(1, ncol + 1):
            cell = ws.cell(row=row, column=c); cell.fill, cell.font = hdr_fill, hdr_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # ---------------- Data sheet (the single source every formula reads)
    wd = wb.active; wd.title = "Data"
    cols = ["week", "store", "date", "item_code", "item_name", "category", "qty", "amount"]
    wd.append(cols); style_header(wd, 1, len(cols))
    for r in data[cols].itertuples(index=False):
        wd.append(list(r))
    n = len(data) + 1
    for i, w in enumerate([12, 11, 12, 11, 30, 12, 7, 11], 1):
        wd.column_dimensions[get_column_letter(i)].width = w
    for row in wd.iter_rows(min_row=2, max_row=n):
        for c in row: c.font = Font(name=F)
        row[2].number_format = "dd-mmm-yy"; row[7].number_format = "#,##0"
    wd.freeze_panes = "A2"; wd.auto_filter.ref = f"A1:H{n}"
    R = lambda col: f"Data!${col}$2:${col}${n}"  # absolute range helper

    # ---------------- Summary sheet
    ws = wb.create_sheet("Summary", 0)
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Kora Home — Weekly Sales Report"; ws["A1"].font = Font(name=F, bold=True, size=16, color="1F3A5F")
    ws["A2"] = f"Week starting {datetime.strptime(week,'%Y-%m-%d'):%d %b %Y}  ·  compared with week of {datetime.strptime(prev_week,'%Y-%m-%d'):%d %b}  ·  sample project, simulated data"
    ws["A2"].font = Font(name=F, italic=True, color="667085")
    ws["A4"], ws["B4"] = "This week", week; ws["A5"], ws["B5"] = "Last week", prev_week
    for c in ("A4", "A5"): ws[c].font = Font(name=F, color="667085")
    for c in ("B4", "B5"): ws[c].font = Font(name=F, color="0000FF")   # inputs: change the week here
    ws["C4"] = "← input cells: change to any week present in Data"; ws["C4"].font = Font(name=F, italic=True, size=9, color="667085")

    ws.append([])
    ws["A7"] = "KPI"; ws["B7"] = "This week"; ws["C7"] = "Last week"; ws["D7"] = "Change"
    style_header(ws, 7, 4)
    kpis = [
        ("Revenue (₹)", f"=SUMIFS({R('H')},{R('A')},$B$4)", f"=SUMIFS({R('H')},{R('A')},$B$5)", "#,##0"),
        ("Units sold", f"=SUMIFS({R('G')},{R('A')},$B$4)", f"=SUMIFS({R('G')},{R('A')},$B$5)", "#,##0"),
        ("Bills (lines)", f"=COUNTIFS({R('A')},$B$4)", f"=COUNTIFS({R('A')},$B$5)", "#,##0"),
        ("Avg value per line (₹)", "=IFERROR(B8/B10,0)", "=IFERROR(C8/C10,0)", "#,##0"),
    ]
    for i, (k, a, b, fmt) in enumerate(kpis, start=8):
        ws.cell(i, 1, k); ws.cell(i, 2, a); ws.cell(i, 3, b); ws.cell(i, 4, f"=IFERROR(B{i}/C{i}-1,0)")
        for c in (2, 3): ws.cell(i, c).number_format = fmt
        ws.cell(i, 4).number_format = '+0.0%;-0.0%;0.0%'
        for c in range(1, 5): ws.cell(i, c).font = Font(name=F, bold=(c == 1)); ws.cell(i, c).border = thin

    ws["A13"] = "By store"; ws["A13"].font = Font(name=F, bold=True, size=12, color="1F3A5F")
    for j, h in enumerate(["Store", "This week (₹)", "Last week (₹)", "Change", "Share of revenue"], 1):
        ws.cell(14, j, h)
    style_header(ws, 14, 5)
    for i, s in enumerate(STORES, start=15):
        ws.cell(i, 1, s)
        ws.cell(i, 2, f"=SUMIFS({R('H')},{R('A')},$B$4,{R('B')},$A{i})")
        ws.cell(i, 3, f"=SUMIFS({R('H')},{R('A')},$B$5,{R('B')},$A{i})")
        ws.cell(i, 4, f"=IFERROR(B{i}/C{i}-1,0)")
        ws.cell(i, 5, f"=IFERROR(B{i}/$B$8,0)")
        ws.cell(i, 2).number_format = ws.cell(i, 3).number_format = "#,##0"
        ws.cell(i, 4).number_format = '+0.0%;-0.0%;0.0%'; ws.cell(i, 5).number_format = "0.0%"
        for c in range(1, 6): ws.cell(i, c).font = Font(name=F); ws.cell(i, c).border = thin
    last_store = 14 + len(STORES)
    red = Font(name=F, color="B42318", bold=True); green = Font(name=F, color="067647", bold=True)
    for rng in (f"D8:D11", f"D15:D{last_store}"):
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["-0.05"], font=red))
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0.05"], font=green))

    cat_start = last_store + 3
    ws.cell(cat_start - 1, 1, "By category").font = Font(name=F, bold=True, size=12, color="1F3A5F")
    for j, h in enumerate(["Category", "This week (₹)", "Last week (₹)", "Change"], 1):
        ws.cell(cat_start, j, h)
    style_header(ws, cat_start, 4)
    cats = sorted(data["category"].dropna().unique())
    for i, c in enumerate(cats, start=cat_start + 1):
        ws.cell(i, 1, c)
        ws.cell(i, 2, f"=SUMIFS({R('H')},{R('A')},$B$4,{R('F')},$A{i})")
        ws.cell(i, 3, f"=SUMIFS({R('H')},{R('A')},$B$5,{R('F')},$A{i})")
        ws.cell(i, 4, f"=IFERROR(B{i}/C{i}-1,0)")
        ws.cell(i, 2).number_format = ws.cell(i, 3).number_format = "#,##0"
        ws.cell(i, 4).number_format = '+0.0%;-0.0%;0.0%'
        for k in range(1, 5): ws.cell(i, k).font = Font(name=F); ws.cell(i, k).border = thin
    cat_end = cat_start + len(cats)
    ws.conditional_formatting.add(f"D{cat_start+1}:D{cat_end}", CellIsRule(operator="lessThan", formula=["-0.05"], font=red))
    ws.conditional_formatting.add(f"D{cat_start+1}:D{cat_end}", CellIsRule(operator="greaterThan", formula=["0.05"], font=green))

    ch = BarChart(); ch.type = "bar"; ch.style = 10; ch.title = "Revenue by store — this week vs last"
    ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=14, max_row=last_store), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=15, max_row=last_store))
    ch.height, ch.width = 7.5, 15
    ws.add_chart(ch, "G7")
    ch2 = BarChart(); ch2.type = "col"; ch2.style = 10; ch2.title = "Revenue by category"
    ch2.add_data(Reference(ws, min_col=2, max_col=3, min_row=cat_start, max_row=cat_end), titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=1, min_row=cat_start + 1, max_row=cat_end))
    ch2.height, ch2.width = 7.5, 15
    for c_, (cur_col, prev_col) in ((ch, ("1F3A5F", "A9BCD3")), (ch2, ("1F3A5F", "A9BCD3"))):
        c_.x_axis.delete = c_.y_axis.delete = False
        c_.legend.position = "b"
    # series order: this week, last week -> dark = this week, light = last week
    for c_ in (ch, ch2):
        c_.series[0].graphicalProperties.solidFill = "1F3A5F"; c_.series[0].graphicalProperties.line.solidFill = "1F3A5F"
        c_.series[1].graphicalProperties.solidFill = "A9BCD3"; c_.series[1].graphicalProperties.line.solidFill = "A9BCD3"
    ws.add_chart(ch2, f"G{cat_start+3}")
    for col, w in zip("ABCDE", [26, 15, 15, 11, 16]):
        ws.column_dimensions[col].width = w
    ws.page_setup.orientation = "landscape"               # prints / PDFs on one page
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1

    # ---------------- Top items (ranked in Python, values via formulas)
    wt = wb.create_sheet("Top Items")
    wt.append(["Rank", "Item code", "Item", "Category", "Units this week", "Revenue this week (₹)", "Revenue last week (₹)", "Change"])
    style_header(wt, 1, 8)
    top = cur.groupby(["item_code", "item_name", "category"])["amount"].sum().nlargest(10).reset_index()
    for i, r in enumerate(top.itertuples(index=False), start=2):
        wt.append([i - 1, r.item_code, r.item_name, r.category,
                   f"=SUMIFS({R('G')},{R('A')},Summary!$B$4,{R('D')},$B{i})",
                   f"=SUMIFS({R('H')},{R('A')},Summary!$B$4,{R('D')},$B{i})",
                   f"=SUMIFS({R('H')},{R('A')},Summary!$B$5,{R('D')},$B{i})",
                   f"=IFERROR(F{i}/G{i}-1,0)"])
        for c in (5, 6, 7): wt.cell(i, c).number_format = "#,##0"
        wt.cell(i, 8).number_format = '+0.0%;-0.0%;0.0%'
    for col, w in zip("ABCDEFGH", [6, 10, 28, 11, 14, 18, 18, 10]):
        wt.column_dimensions[col].width = w
    wt.cell(13, 1, "Ranking is fixed for the report week; values recalculate from the Data sheet.").font = Font(name=F, italic=True, size=9, color="667085")

    # ---------------- Reorder alerts
    wr = wb.create_sheet("Reorder Alerts")
    wr["A1"] = "Reorder threshold (days of cover):"; wr["D1"] = REORDER_DAYS; wr["D1"].font = Font(name=F, color="0000FF", bold=True)
    wr["D1"].fill = PatternFill("solid", fgColor="FFFF00")
    wr.append([])
    wr.append(["Store", "Item code", "Item", "Closing stock", "Sold this week", "Avg daily sales", "Days of cover", "Status"])
    style_header(wr, 3, 8)
    st = stock.merge(cur.groupby(["store", "item_code"])["qty"].sum().reset_index(), on=["store", "item_code"], how="left") \
              .merge(pd.read_csv(IN / "item_master.csv")[["item_code", "item_name"]], on="item_code")
    st["qty"] = st["qty"].fillna(0)
    st["cover"] = st["closing_stock"] / (st["qty"] / 7).where(st["qty"] > 0)
    st = st.sort_values("cover", na_position="last")
    for i, r in enumerate(st.itertuples(index=False), start=4):
        wr.append([r.store, r.item_code, r.item_name, r.closing_stock,
                   f"=SUMIFS({R('G')},{R('A')},Summary!$B$4,{R('B')},$A{i},{R('D')},$B{i})",
                   f"=E{i}/7", f'=IF(F{i}=0,"no sales",D{i}/F{i})',
                   f'=IF(D{i}=0,"OUT OF STOCK",IF(F{i}=0,"OK",IF(G{i}<$D$1,"REORDER NOW","OK")))'])
        wr.cell(i, 6).number_format = "0.0"; wr.cell(i, 7).number_format = "0.0"
    last = 3 + len(st)
    wr.conditional_formatting.add(f"H4:H{last}", CellIsRule(operator="equal", formula=['"OUT OF STOCK"'], font=red,
                                  fill=PatternFill("solid", fgColor="FEE4E2")))
    wr.conditional_formatting.add(f"H4:H{last}", CellIsRule(operator="equal", formula=['"REORDER NOW"'],
                                  font=Font(name=F, color="B54708", bold=True), fill=PatternFill("solid", fgColor="FEF0C7")))
    for col, w in zip("ABCDEFGH", [10, 10, 28, 13, 14, 15, 13, 16]):
        wr.column_dimensions[col].width = w
    wr.freeze_panes = "A4"; wr.auto_filter.ref = f"A3:H{last}"

    # ---------------- Data log
    wl = wb.create_sheet("Data Log")
    wl.append(["Source", "Issue found", "Rows", "What the script did"]); style_header(wl, 1, 4)
    for s, iss, rows, act in LOG:
        wl.append([s, iss, rows if rows else "all", act])
    for col, w in zip("ABCD", [12, 52, 8, 52]):
        wl.column_dimensions[col].width = w

    for sh in wb.worksheets:
        for row in sh.iter_rows():
            for c in row:
                if c.font and c.font.name != F:
                    c.font = Font(name=F, bold=c.font.bold, italic=c.font.italic, size=c.font.size, color=c.font.color)

    xlsx = OUT / f"Kora_Home_Weekly_Report_{week}.xlsx"
    wb.save(xlsx)

    # ---------------- E-mail summary
    rev, prev_rev = cur.amount.sum(), prev.amount.sum()
    by_store = cur.groupby("store").amount.sum() / prev.groupby("store").amount.sum() - 1
    by_cat = cur.groupby("category").amount.sum() / prev.groupby("category").amount.sum() - 1
    oos = st[st.closing_stock == 0]
    low = st[(st.closing_stock > 0) & (st.cover < REORDER_DAYS)]
    best = top.iloc[0]
    html = f"""<div style="font-family:Arial,sans-serif;max-width:620px;color:#101828">
<h2 style="color:#1F3A5F;margin-bottom:4px">Kora Home · Week of {datetime.strptime(week,'%Y-%m-%d'):%d %b %Y}</h2>
<p style="color:#667085;margin-top:0">Automated weekly summary · sample project, simulated data</p>
<p><b>Revenue ₹{rev:,.0f}</b> ({rev/prev_rev-1:+.1%} vs last week) from {int(cur.qty.sum()):,} units.</p>
<ul>
<li><b>Fastest growing category:</b> {by_cat.idxmax()} ({by_cat.max():+.0%}); weakest: {by_cat.idxmin()} ({by_cat.min():+.0%}).</li>
<li><b>Store to watch:</b> {by_store.idxmin()} ({by_store.min():+.1%}) {'— the only store that declined' if (by_store < 0).sum() == 1 else '— lowest growth of all stores'}.</li>
<li><b>Best seller:</b> {best.item_name} (₹{best.amount:,.0f}).</li>
<li><b>Stock:</b> {len(oos)} item-store {'combination is' if len(oos)==1 else 'combinations are'} out of stock and {len(low)} more have under {REORDER_DAYS} days of cover — see the Reorder Alerts sheet.</li>
</ul>
<p style="color:#667085;font-size:12px">Full report attached: {xlsx.name}</p></div>"""
    (OUT / f"email_summary_{week}.html").write_text(html, encoding="utf-8")
    print(f"Saved {xlsx.name} and email summary. Revenue {rev:,.0f} vs {prev_rev:,.0f}; OOS {len(oos)}; low {len(low)}")
    print(by_store.round(3).to_dict(), by_cat.round(3).to_dict())

if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--week", required=True, help="Monday of the report week, YYYY-MM-DD")
    main(ap.parse_args().week)
